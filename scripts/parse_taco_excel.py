import os
import sys
import json
import hashlib
import re
import unicodedata
from datetime import datetime, timezone
import openpyxl

RAW_XLSX_PATH = 'src/data/taco/raw/taco_nepa_oficial.xlsx'
OUTPUT_JSON_PATH = 'src/data/taco/taco_4_edicao_complete.json'
OFFICIAL_NEPA_URL = 'https://www.nepa.unicamp.br/arquivo/uploads/taco-4a-edicao/taco-4a-edicao-2/'
EXPECTED_NEPA_SHA256 = 'A66B8EC528DAEABC63BC2B015FC9BD8C6D76B941C2FC0ED93A4311D449302D14'
EXPECTED_FILE_SIZE = 322270

# Mapeamento oficial de colunas de nutrientes na aba 'CMVCol taco3'
# As colunas em openpyxl são 1-indexed (coluna 1 = Código, coluna 2 = Descrição)
NUTRIENT_COL_MAP = {
    3: ('moisture_pct', 'Umidade', '%'),
    4: ('energy_kcal', 'Energia', 'kcal'),
    5: ('energy_kj', 'Energia', 'kJ'),
    6: ('protein_g', 'Proteína', 'g'),
    7: ('fat_g', 'Lipídeos', 'g'),
    8: ('cholesterol_mg', 'Colesterol', 'mg'),
    9: ('carbohydrate_g', 'Carboidrato', 'g'),
    10: ('fiber_g', 'Fibra Alimentar', 'g'),
    11: ('ash_g', 'Cinzas', 'g'),
    12: ('calcium_mg', 'Cálcio', 'mg'),
    13: ('magnesium_mg', 'Magnésio', 'mg'),
    14: ('manganese_mg', 'Manganês', 'mg'),
    15: ('phosphorus_mg', 'Fósforo', 'mg'),
    16: ('iron_mg', 'Ferro', 'mg'),
    17: ('sodium_mg', 'Sódio', 'mg'),
    18: ('potassium_mg', 'Potássio', 'mg'),
    19: ('copper_mg', 'Cobre', 'mg'),
    20: ('zinc_mg', 'Zinco', 'mg'),
    21: ('retinol_mcg', 'Retinol', 'mcg'),
    22: ('re_mcg', 'Equivalente de Retinol (RE)', 'mcg'),
    23: ('rae_mcg', 'Equivalente de Atividade de Retinol (RAE)', 'mcg'),
    24: ('thiamine_mg', 'Tiamina', 'mg'),
    25: ('riboflavin_mg', 'Riboflavina', 'mg'),
    26: ('pyridoxine_mg', 'Piridoxina', 'mg'),
    27: ('niacin_mg', 'Niacina', 'mg'),
    28: ('vitamin_c_mg', 'Vitamina C', 'mg'),
}

KNOWN_GROUPS = [
    'Cereais e derivados',
    'Verduras, hortaliças e derivados',
    'Frutas e derivados',
    'Gorduras e óleos',
    'Pescados e frutos do mar',
    'Carnes e derivados',
    'Leite e derivados',
    'Bebidas (alcoólicas e não alcoólicas)',
    'Ovos e derivados',
    'Produtos açucarados',
    'Miscelâneas',
    'Outros alimentos industrializados',
    'Alimentos preparados',
    'Leguminosas e derivados',
    'Nozes e sementes'
]

# Registro canônico e auditável de reconciliações versionadas
# Nenhuma correção silenciosa: cada transformação possui código, valor raw de cada aba, valor canônico e justificativa
RECONCILIATION_REGISTRY = [
    {
        'rule_id': 'RECON_TACO4_FOOD_540_NAME',
        'food_code': '540',
        'raw_cmv_value': 'L',
        'raw_ag_value': 'Feijoada',
        'selected_canonical_value': 'Feijoada',
        'source': 'TACO 4ª edição oficial NEPA/UNICAMP',
        'version': '4.0.0',
        'reason': 'Aba CMVCol taco3 contém caractere corrompido/truncado "L" na linha do alimento 540. A aba AGtaco3 da mesma planilha oficial contém a descrição completa e correta "Feijoada". Reconciliação canônica auditável intra-arquivo.'
    }
]

def compute_sha256(filepath):
    h = hashlib.sha256()
    with open(filepath, 'rb') as f:
        while chunk := f.read(65536):
            h.update(chunk)
    return h.hexdigest().upper()

def normalize_text(text):
    if not text:
        return ''
    nfkd = unicodedata.normalize('NFKD', text)
    no_accents = ''.join([c for c in nfkd if not unicodedata.combining(c)])
    clean = re.sub(r'[^a-zA-Z0-9\s]', ' ', no_accents).lower()
    return re.sub(r'\s+', ' ', clean).strip()

def detect_preparation_state(name, food_group):
    name_lower = name.lower()
    if 'cozido' in name_lower or 'cozida' in name_lower:
        return 'cooked'
    if 'grelhado' in name_lower or 'grelhada' in name_lower:
        return 'grilled'
    if 'assado' in name_lower or 'assada' in name_lower:
        return 'roasted'
    if 'frito' in name_lower or 'frita' in name_lower:
        return 'fried'
    if 'cru' in name_lower or 'crua' in name_lower:
        return 'raw'
    if food_group == 'Alimentos preparados':
        return 'prepared'
    if food_group == 'Outros alimentos industrializados':
        return 'industrialized'
    return 'other'

def parse_cell_value(val):
    """
    Interpreta o valor analítico rigorosamente idêntico ao parser TypeScript:
    - 'Tr' / 'tr' -> TRACE (Traço analítico detectado abaixo do limite de quantificação)
    - 'NA' / 'na' -> NOT_APPLICABLE (Não Aplicável para aquela matriz alimentar)
    - '*' -> UNDER_REEVALUATION (Análise sob reavaliação metodológica)
    - '' (blank) -> NOT_REQUESTED (Análise não solicitada)
    - Números negativos -> UNKNOWN (Rejeitados explicitamente, nunca convertidos em zero)
    - ',5' ou '.5' -> 0.5 (Não remove pontuação decimal inicial)
    - ',0,02' -> 0.02 (Correção de typo de pontuação dupla conhecida da TACO)
    - '12abc' ou lixo -> UNKNOWN (Validação estrita de formato numérico)
    - '0' / 0.0 -> KNOWN_NUMERIC_ZERO (Zero laboratorial comprovado)
    - número > 0 -> NUMERIC_VALUE
    """
    if val is None:
        return {
            'raw_value': '',
            'numeric_value': None,
            'value_status': 'NOT_REQUESTED',
            'data_quality': 'not_requested',
        }
    
    raw_str = str(val).strip()

    if raw_str in ['Tr', 'tr']:
        return {
            'raw_value': 'Tr',
            'numeric_value': None,
            'value_status': 'TRACE',
            'data_quality': 'trace',
        }
    if raw_str == 'NA' or raw_str.lower() == 'na':
        return {
            'raw_value': 'NA',
            'numeric_value': None,
            'value_status': 'NOT_APPLICABLE',
            'data_quality': 'not_applicable',
        }
    if raw_str == '*':
        return {
            'raw_value': '*',
            'numeric_value': None,
            'value_status': 'UNDER_REEVALUATION',
            'data_quality': 'under_reevaluation',
        }
    if raw_str == '':
        return {
            'raw_value': '',
            'numeric_value': None,
            'value_status': 'NOT_REQUESTED',
            'data_quality': 'not_requested',
        }

    # Rejeição explícita de números negativos (não podem virar KNOWN_NUMERIC_ZERO)
    if raw_str.startswith('-'):
        return {
            'raw_value': raw_str,
            'numeric_value': None,
            'value_status': 'UNKNOWN',
            'data_quality': 'unknown',
        }

    # Normalização determinística de decimais
    normalized_str = raw_str
    if normalized_str == ',0,02':
        normalized_str = '0.02'
    elif normalized_str.startswith(',') or normalized_str.startswith('.'):
        normalized_str = '0.' + normalized_str[1:]
    else:
        normalized_str = normalized_str.replace(',', '.')

    # Validação estrita: rejeita strings com sufixos ou caracteres inválidos (ex: "12abc")
    if not re.match(r'^\d+(\.\d+)?$', normalized_str):
        return {
            'raw_value': raw_str,
            'numeric_value': None,
            'value_status': 'UNKNOWN',
            'data_quality': 'unknown',
        }

    try:
        f = float(normalized_str)
        rounded = round(f, 4)
        if rounded == 0.0:
            return {
                'raw_value': raw_str,
                'numeric_value': 0.0,
                'value_status': 'KNOWN_NUMERIC_ZERO',
                'data_quality': 'analytical',
            }
        return {
            'raw_value': raw_str,
            'numeric_value': rounded,
            'value_status': 'NUMERIC_VALUE',
            'data_quality': 'analytical',
        }
    except ValueError:
        return {
            'raw_value': raw_str,
            'numeric_value': None,
            'value_status': 'UNKNOWN',
            'data_quality': 'unknown',
        }

def evaluate_atwater_consistency(energy_kcal, protein_g, carbs_g, fat_g):
    if energy_kcal is None or protein_g is None or carbs_g is None or fat_g is None:
        return 'insufficient_data'
    
    calc_kcal = round((4.0 * protein_g + 4.0 * carbs_g + 9.0 * fat_g) * 10) / 10.0
    if energy_kcal == 0 and calc_kcal == 0:
        return 'consistent'
    if energy_kcal == 0:
        return 'different_from_macro_estimate'
    
    delta = abs(energy_kcal - calc_kcal)
    delta_pct = round((delta / energy_kcal) * 1000) / 10.0
    if delta_pct <= 10.0:
        return 'consistent'
    if delta_pct <= 20.0:
        return 'within_tolerance'
    return 'different_from_macro_estimate'

def main():
    print('1. Verificando presença e integridade do arquivo TACO Excel oficial do NEPA/UNICAMP...')
    if not os.path.exists(RAW_XLSX_PATH):
        raise FileNotFoundError(
            f'Arquivo oficial não encontrado em: {RAW_XLSX_PATH}\n'
            f'O arquivo oficial da TACO deve ser obtido diretamente do NEPA/UNICAMP:\n'
            f'URL oficial: {OFFICIAL_NEPA_URL}'
        )

    file_size = os.path.getsize(RAW_XLSX_PATH)
    file_checksum = compute_sha256(RAW_XLSX_PATH)
    print(f'Tamanho do arquivo:  {file_size} bytes (esperado: {EXPECTED_FILE_SIZE})')
    print(f'Checksum calculado: {file_checksum}')
    print(f'Checksum esperado:   {EXPECTED_NEPA_SHA256}')
    
    if file_checksum != EXPECTED_NEPA_SHA256:
        raise RuntimeError(
            f'Checksum do arquivo Excel ({file_checksum}) não corresponde ao oficial do NEPA ({EXPECTED_NEPA_SHA256})!'
        )

    print('2. Carregando planilha oficial do NEPA/UNICAMP (openpyxl)...')
    wb = openpyxl.load_workbook(RAW_XLSX_PATH, data_only=True)
    ws_cmv = wb['CMVCol taco3']
    ws_ag = wb['AGtaco3']

    # Mapa auxiliar de nomes da aba AGtaco3 para desambiguação e auditoria
    ag_names = {}
    for r in range(1, ws_ag.max_row + 1):
        c0 = ws_ag.cell(r, 1).value
        c1 = ws_ag.cell(r, 2).value
        if isinstance(c0, int):
            ag_names[c0] = str(c1).strip() if c1 else ''

    current_group = None
    foods = []
    stats = {
        'total_rows_source': 597,
        'imported': 0,
        'eligible_for_engine': 0,
        'incomplete_nutrition': 0,
        'consistent_energy': 0,
        'within_tolerance_energy': 0,
        'divergent_energy': 0,
        'insufficient_data_energy': 0,
    }

    def norm_str(s):
        return ''.join(c for c in unicodedata.normalize('NFKD', s) if not unicodedata.combining(c)).lower()

    print('3. Processando os 597 alimentos oficiais e seus grupos canônicos...')
    for r in range(1, ws_cmv.max_row + 1):
        val0 = ws_cmv.cell(r, 1).value
        val1 = ws_cmv.cell(r, 2).value

        # Identificação de grupos alimentares no cabeçalho das seções da planilha oficial
        if isinstance(val0, str):
            v_norm = norm_str(val0.strip())
            for g in KNOWN_GROUPS:
                if norm_str(g) == v_norm:
                    current_group = g
                    break
            continue

        if not isinstance(val0, int) or val0 < 1 or val0 > 597:
            continue

        food_code_num = val0
        code_str = f'{food_code_num:03d}'
        food_name = str(val1).strip() if val1 else ''

        # Reconciliação canônica auditável: registro formal e explícito
        reconciliation_info = None
        if food_code_num == 540:
            recon = next(item for item in RECONCILIATION_REGISTRY if item['food_code'] == '540')
            food_name = recon['selected_canonical_value']
            reconciliation_info = recon

        if current_group is None:
            raise ValueError(f'Alimento {food_code_num} sem grupo alimentar associado!')

        food_group = current_group
        prep_state = detect_preparation_state(food_name, food_group)

        # Extração de nome científico se presente (ex: "Oryza sativa L.")
        sci_match = re.search(r'\((([A-Z][a-z]+)\s+([a-z]+(?:\s+var\.\s+[a-z]+)?))\)', food_name)
        scientific_name = sci_match.group(1) if sci_match else None

        nutrients = {}
        for col_idx, (nut_code, nut_name, nut_unit) in NUTRIENT_COL_MAP.items():
            cell_val = ws_cmv.cell(r, col_idx).value
            parsed = parse_cell_value(cell_val)
            nutrients[nut_code] = {
                'code': nut_code,
                'name': nut_name,
                'unit': nut_unit,
                'raw_value': parsed['raw_value'],
                'numeric_value': parsed['numeric_value'],
                'amount_per_100g': parsed['numeric_value'],
                'value_status': parsed['value_status'],
                'data_quality': parsed['data_quality'],
                'source_reference': f'TACO 4ª edição, alimento {code_str}, col {col_idx}',
            }

        # Validações obrigatórias de macronutrientes para motor determinístico
        e_kcal = nutrients.get('energy_kcal', {}).get('numeric_value')
        p_g = nutrients.get('protein_g', {}).get('numeric_value')
        c_g = nutrients.get('carbohydrate_g', {}).get('numeric_value')
        f_g = nutrients.get('fat_g', {}).get('numeric_value')

        if e_kcal is not None and p_g is not None and c_g is not None and f_g is not None:
            eligibility = 'eligible_for_engine'
            stats['eligible_for_engine'] += 1
        else:
            eligibility = 'incomplete_nutrition'
            stats['incomplete_nutrition'] += 1

        energy_status = evaluate_atwater_consistency(e_kcal, p_g, c_g, f_g)
        if energy_status == 'consistent':
            stats['consistent_energy'] += 1
        elif energy_status == 'within_tolerance':
            stats['within_tolerance_energy'] += 1
        elif energy_status == 'different_from_macro_estimate':
            stats['divergent_energy'] += 1
        else:
            stats['insufficient_data_energy'] += 1

        food_entry = {
            'source_food_code': code_str,
            'name': food_name,
            'normalized_name': normalize_text(food_name),
            'scientific_name': scientific_name,
            'food_group': food_group,
            'preparation_state': prep_state,
            'is_generic': True,
            'is_active': True,
            'source_type': 'official',
            'validation_status': 'official_approved',
            'engine_eligibility_status': eligibility,
            'energy_consistency_status': energy_status,
            'nutrients': nutrients,
        }
        if reconciliation_info:
            food_entry['reconciliation'] = reconciliation_info

        foods.append(food_entry)
        stats['imported'] += 1

    print(f'Total de alimentos processados com sucesso: {len(foods)}')
    if len(foods) != 597:
        raise RuntimeError(f'Esperava exatamente 597 alimentos oficiais, obteve {len(foods)}')

    # Metadados determinísticos para cálculo do checksum do dataset normalizado
    # Exclui timestamps variáveis em tempo de execução para garantir estrita reprodutibilidade
    deterministic_source_metadata = {
        'id': 'taco_4_edicao',
        'name': 'Tabela Brasileira de Composição de Alimentos - TACO',
        'short_name': 'TACO',
        'version': '4.0.0',
        'publisher': 'NEPA - Núcleo de Estudos e Pesquisas em Alimentação / UNICAMP',
        'domain': 'nepa.unicamp.br',
        'source_url': OFFICIAL_NEPA_URL,
        'reference': 'NEPA/UNICAMP. Tabela brasileira de composição de alimentos - TACO. 4. ed. rev. e ampl. Campinas: NEPA-UNICAMP, 2011. 161 p.',
        'source_file': 'taco_nepa_oficial.xlsx',
        'source_file_checksum': file_checksum,
        'source_file_size_bytes': file_size,
        'parser_version': '2.1.0',
        'normalization_version': '2.1.0',
        'license_notes': 'Uso institucional, científico e acadêmico autorizado com citação obrigatória da fonte NEPA/UNICAMP.',
        'reconciliations': RECONCILIATION_REGISTRY,
        'stats': stats,
    }

    # Cálculo determinístico do checksum do dataset normalizado
    # O hash é calculado exclusivamente sobre o conteúdo estático e ordenado dos alimentos e metadados determinísticos
    deterministic_payload = {
        'source': deterministic_source_metadata,
        'foods': foods,
    }
    deterministic_serialized = json.dumps(deterministic_payload, sort_keys=True, ensure_ascii=False).encode('utf-8')
    dataset_checksum = hashlib.sha256(deterministic_serialized).hexdigest().upper()
    deterministic_source_metadata['dataset_checksum'] = dataset_checksum

    # Montagem do artefato canônico com metadata operacional claramente separada
    canonical_output = {
        'source': deterministic_source_metadata,
        'operational_metadata': {
            'import_run_id': 'canonical_nepa_import_20260908',
            'imported_at': '2026-09-08T00:00:00Z',
            'execution_timestamp': datetime.now(timezone.utc).isoformat(),
        },
        'foods': foods,
    }

    with open(OUTPUT_JSON_PATH, 'w', encoding='utf-8') as f:
        json.dump(canonical_output, f, ensure_ascii=False, indent=2)

    output_size = os.path.getsize(OUTPUT_JSON_PATH)
    output_checksum = compute_sha256(OUTPUT_JSON_PATH)
    print(f'4. Artefato canônico exportado com sucesso:')
    print(f'   Arquivo:  {OUTPUT_JSON_PATH}')
    print(f'   Tamanho:  {output_size} bytes')
    print(f'   Dataset Checksum: {dataset_checksum}')
    print(f'   File Checksum:    {output_checksum}')
    print('Estatísticas oficiais finais:')
    for k, v in stats.items():
        print(f'  - {k}: {v}')

if __name__ == '__main__':
    main()
